from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime
import json
from typing import Any, Dict, List


def set_cell_background(cell, fill):
    """Set cell background color"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), fill)
    cell._element.get_or_add_tcPr().append(shading_elm)


def flatten_json(data: Any, parent_key: str = '', sep: str = '.') -> Dict:
    """Flatten nested JSON structure"""
    items = []
    if isinstance(data, dict):
        for k, v in data.items():
            new_key = f"{parent_key}{sep}{k}" if parent_key else k
            if isinstance(v, (dict, list)):
                items.extend(flatten_json(v, new_key, sep=sep).items())
            else:
                items.append((new_key, v))
    elif isinstance(data, list):
        for i, v in enumerate(data):
            new_key = f"{parent_key}[{i}]"
            if isinstance(v, (dict, list)):
                items.extend(flatten_json(v, new_key, sep=sep).items())
            else:
                items.append((new_key, v))
    else:
        return {parent_key: data}
    return dict(items)


def json_to_table_data(data: Any) -> List[List[str]]:
    """Convert JSON to table format"""
    if isinstance(data, dict):
        flattened = flatten_json(data)
        rows = [['Field', 'Value']]
        for key, value in flattened.items():
            rows.append([str(key), str(value)[:100]])  # Limit value length
        return rows
    elif isinstance(data, list) and len(data) > 0:
        if isinstance(data[0], dict):
            # List of dictionaries
            keys = set()
            for item in data:
                if isinstance(item, dict):
                    keys.update(item.keys())
            keys = sorted(list(keys))
            rows = [[key for key in keys]]
            for item in data:
                if isinstance(item, dict):
                    rows.append([str(item.get(k, '')) for k in keys])
            return rows
    return [['Data', str(data)[:100]]]


def add_table_to_doc(doc, title: str, data: List[List[str]]):
    """Add a formatted table to the document"""
    doc.add_heading(title, level=3)

    if not data or len(data) == 0:
        doc.add_paragraph("No data available")
        return

    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.style = 'Light Grid Accent 1'

    # Style header row
    for i, cell in enumerate(table.rows[0].cells):
        cell.text = str(data[0][i])
        set_cell_background(cell, 'D3D3D3')
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    # Add data rows
    for row_idx, row_data in enumerate(data[1:], 1):
        for col_idx, cell_value in enumerate(row_data):
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = str(cell_value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)


def generate_workflow_report(
    project_name: str,
    workflow_name: str,
    execution_data: Dict[str, Any],
    agent_outputs: Dict[str, Any]
) -> Document:
    """
    Generate a Word document report from workflow execution data

    Args:
        project_name: Name of the project
        workflow_name: Name of the workflow
        execution_data: Main workflow execution data
        agent_outputs: Dictionary of agent outputs {agent_name: output_data}
    """
    doc = Document()

    # Add title
    title = doc.add_heading('Workflow Execution Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add metadata
    doc.add_heading('Executive Summary', level=1)
    metadata_table_data = [
        ['Field', 'Value'],
        ['Project Name', project_name],
        ['Workflow Name', workflow_name],
        ['Generated On', datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
        ['Status', execution_data.get('status', 'Unknown')]
    ]
    add_table_to_doc(doc, 'Execution Metadata', metadata_table_data)

    # Add execution summary
    doc.add_heading('Execution Details', level=1)
    execution_summary = execution_data.get('summary', {})
    if execution_summary:
        summary_table = json_to_table_data(execution_summary)
        add_table_to_doc(doc, 'Summary', summary_table)

    # Add agent outputs
    doc.add_heading('Agent Outputs', level=1)

    if agent_outputs:
        for agent_name, agent_data in agent_outputs.items():
            doc.add_heading(f'{agent_name} Output', level=2)

            if isinstance(agent_data, dict):
                table_data = json_to_table_data(agent_data)
                add_table_to_doc(doc, agent_name, table_data)
            elif isinstance(agent_data, list):
                table_data = json_to_table_data(agent_data)
                add_table_to_doc(doc, agent_name, table_data)
            else:
                doc.add_paragraph(str(agent_data))
    else:
        doc.add_paragraph("No agent outputs available")

    # Add footer
    doc.add_paragraph()
    footer_para = doc.add_paragraph('This report was automatically generated by the Academian Agentic Platform')
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer_para.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(128, 128, 128)

    return doc


def save_report(doc: Document, filepath: str):
    """Save the document to a file"""
    doc.save(filepath)
